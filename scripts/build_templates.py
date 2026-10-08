from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
import sys

BASE=Path(__file__).resolve().parents[1]
OUT=BASE/'templates/sheets'
OUT.mkdir(parents=True,exist_ok=True)
if list(OUT.glob('*.xlsx')) and '--force' not in sys.argv:
    raise SystemExit('Workbook templates already exist; inspect them first, then pass --force to replace them.')
BLUE='245B70'
def workbook(path,sheets):
    wb=Workbook(); wb.remove(wb.active)
    for name,headers,notes,validations in sheets:
        ws=wb.create_sheet(name)
        ws.append(headers)
        for c in ws[1]:
            c.font=Font(bold=True,color='FFFFFF')
            c.fill=PatternFill('solid',fgColor=BLUE)
            c.alignment=Alignment(wrap_text=True,vertical='top')
        ws.freeze_panes='A2'
        if headers: ws.auto_filter.ref=f'A1:{ws.cell(1,len(headers)).column_letter}200'
        for i,h in enumerate(headers,1):
            ws.column_dimensions[ws.cell(1,i).column_letter].width=min(max(len(h)*2,14),32)
        ws.row_dimensions[1].height=36
        for note in notes: ws.append([note])
        for formula,range_ref in validations:
            dv=DataValidation(type='list',formula1=formula,allow_blank=True)
            ws.add_data_validation(dv); dv.add(range_ref)
    wb.save(path)

workbook(OUT/'校務行動追蹤範本.xlsx',[
('ActionTracker',['編號','工作項目','來源／決議','承辦角色','期限','狀態','最後更新','待確認／阻礙'],['一列一項行動；討論意見不可填成正式決議。','狀態：未開始、進行中、完成、待確認、暫停。'],[('"未開始,進行中,完成,待確認,暫停"','F2:F200')]),
('SourceLog',['來源ID','檔名','版本／日期','來源類型','核定狀態','使用欄位／摘錄','疑點'],['記錄每筆追蹤資料的依據。'],[]),
('ChangeLog',['日期','操作者角色','變更欄位','變更前','變更後','原因／來源'],['每次清理或狀態更新留下紀錄。'],[])
])
workbook(OUT/'問卷品質檢查範本.xlsx',[
('Raw',['response_id','submitted_at','group','Q1','Q2','Q3','comment'],['貼入原始匯出副本；不要覆寫來源檔。'],[]),
('Cleaned',['response_id','submitted_at','group','Q1_clean','Q2_clean','Q3_clean','include_flag','exclusion_reason'],['清理規則需先確認；保留原始值供比對。'],[('"是,否,待確認"','G2:G200')]),
('CleaningLog',['issue_id','row_id','issue_type','proposed_rule','approved_by_role','decision','notes'],['列出重複、缺值、混合尺度和標籤不一致。'],[]),
('Summary',['metric','numerator','denominator','result','period','source_sheet','limitation'],['每個比例寫分子、分母、期間和限制。'],[])
])
workbook(OUT/'課程意向與容量規劃範本.xlsx',[
('InterestRaw',['response_id','grade_group','choice_1','choice_2','submitted_at','source_version'],['意向調查不是正式選課結果。'],[]),
('CourseCapacity',['course_name','proposed_min','proposed_max','room_type','teacher_capacity','schedule_constraint','approval_status'],['容量及開課狀態需由教務角色確認。'],[('"草案,待確認,核定"','G2:G200')]),
('Reconciliation',['course_label_raw','canonical_label_candidate','mapping_source','decision_owner','decision_status'],['舊課名映射不得由 Agent 自行合併。'],[('"待確認,已確認,排除"','E2:E200')])
])
workbook(OUT/'設備盤點與差異範本.xlsx',[
('Inventory',['asset_id','item','location','book_qty','physical_qty','condition','threshold','count_date','difference','review_status'],['盤點差異是待複核訊號，不是採購核准。'],[('"未盤點,待複核,已核對"','J2:J200')]),
('Exceptions',['asset_id','issue','source_row','next_check','owner_role','decision_status'],['只列差異、重複、缺漏和現場待查。'],[])
])
print('Created',len(list(OUT.glob('*.xlsx'))),'workbooks in',OUT)
